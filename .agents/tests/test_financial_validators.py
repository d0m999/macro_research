from __future__ import annotations

import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock


REPOSITORY_ROOT = Path(__file__).resolve().parents[2]


def load_module(name: str, relative_path: str):
    spec = importlib.util.spec_from_file_location(name, REPOSITORY_ROOT / relative_path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Unable to load {relative_path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


artifact_validator = load_module(
    "validate_financial_artifact",
    ".agents/scripts/validate_financial_artifact.py",
)


try:
    import openpyxl
    from openpyxl.worksheet.table import Table
except ImportError:
    openpyxl = None


class FormulaLiteralTests(unittest.TestCase):
    def test_error_markers_inside_string_literals_are_ignored(self) -> None:
        formulas = (
            '=COUNTIF(A1:A3,"#REF!")',
            '=IFERROR(A1,"#N/A")',
            '="escaped ""#VALUE!"" text"',
            "='FY\"26'!A1&\"#REF!\"",
            "='FY''\"26'!A1&\"escaped \"\"#REF!\"\" text\"",
            "='FY#REF!\"26'!A1&\"it's #REF!\"",
            '=SUM(Table1[\'#Data])&"#REF!"',
            '=SUM(Table1[Column\']name])&"#REF!"',
            '=SUM(Table1[Column\'\'name])&"#REF!"',
            '=SUM(Table1[[Check "\'#REF!"]])',
            '=SUM(Table1[[Check "\'#VALUE!"]])',
            '=SUM(Table1[[Check "\'#DIV/0!"]])',
            '=SUM(Table1[[Check "\'#NAME?"]])',
            '=SUM(Table1[[Check "\'#NULL!"]])',
            '=SUM(Table1[[Check "\'#NUM!"]])',
            '=SUM(Table1[[Check "\'#N/A"]])',
        )
        for formula in formulas:
            with self.subTest(formula=formula):
                self.assertFalse(artifact_validator.formula_has_unquoted_excel_error(formula))

    def test_error_markers_outside_string_literals_are_rejected(self) -> None:
        formulas = (
            "=#REF!",
            '=IFERROR(#N/A,"fallback")',
            '=IF(A1="#REF!",#VALUE!,0)',
            '=IF(A1,"unterminated #REF!)',
            "='FY\"26'!A1+#REF!&\" USD\"",
            "='FY''\"26'!A1+#REF!&\" USD\"",
            "=IFERROR('FY\"26'!#REF!,\"#N/A\")",
            "='FY\"26'!A1+'Q\"27'!#REF!&\" USD\"",
            "='[FY\"26.xlsx]Budget''s'!A1+#REF!&\" USD\"",
            '=Table1[\'#Data]+#REF!+Table1[\'#Totals]&" USD"',
            '=Table1[Column\']name]+#REF!+Table1[Column\'\'name]&" USD"',
            "='unterminated #REF!",
            '=SUM(Table1[[Check "\'#REF!"]])+#REF!',
            '=#REF!+SUM(Table1[[Check "\'#REF!"]])',
        )
        for formula in formulas:
            with self.subTest(formula=formula):
                self.assertTrue(artifact_validator.formula_has_unquoted_excel_error(formula))


@unittest.skipUnless(openpyxl, "openpyxl is required for workbook validation tests")
class WorkbookValidationTests(unittest.TestCase):
    def test_artifact_validator_checks_formula_literals_without_cached_results(self) -> None:
        cases = (
            (None, ('=COUNTIF(B1:B3,"#REF!")', '=IFERROR(B2,"#N/A")'), 0),
            ('FY"26', ("='FY\"26'!A1+#REF!&\" USD\"",), 1),
            ('FY\'"26', ("='FY''\"26'!A1+#REF!&\" USD\"",), 1),
            ('FY"26', ("='FY\"26'!A1&\"#REF!\"",), 0),
            ('FY\'"26', ("='FY''\"26'!A1&\"escaped \"\"#REF!\"\" text\"",), 0),
        )
        for sheet_name, formulas, expected_exit_code in cases:
            with self.subTest(sheet_name=sheet_name, formulas=formulas):
                with tempfile.TemporaryDirectory() as directory:
                    path = Path(directory) / "quoted-errors.xlsx"
                    workbook = openpyxl.Workbook()
                    for row, formula in enumerate(formulas, start=1):
                        workbook.active.cell(row, 1, formula)
                    if sheet_name is not None:
                        workbook.create_sheet(sheet_name)["A1"] = 1
                    workbook.save(path)

                    cached_workbook = openpyxl.load_workbook(path, data_only=True)
                    try:
                        for row in range(1, len(formulas) + 1):
                            self.assertIsNone(cached_workbook.active.cell(row, 1).value)
                    finally:
                        cached_workbook.close()
                    result, exit_code = artifact_validator.validate(path)

                self.assertEqual(exit_code, expected_exit_code)
                self.assertEqual(result["status"], "fail" if expected_exit_code else "pass_with_limitations")
                self.assertEqual(result["structural_checks"]["formula_text"], "fail" if expected_exit_code else "pass")

    def test_cli_distinguishes_escaped_table_headers_from_actual_errors(self) -> None:
        reference = 'Table1[[Check "\'#REF!"]]'
        cases = (
            (f'=SUM({reference})', 0),
            (f'=IFERROR(SUM({reference}),"#REF!")', 0),
            (f'=SUM({reference})+#REF!', 1),
            (f'=#REF!+SUM({reference})', 1),
        )
        for formula, expected_exit_code in cases:
            with self.subTest(formula=formula):
                with tempfile.TemporaryDirectory() as directory:
                    path = Path(directory) / "structured-reference.xlsx"
                    workbook = openpyxl.Workbook()
                    sheet = workbook.active
                    sheet["A1"] = 'Check "#REF!"'
                    sheet["A2"] = 123
                    sheet.add_table(Table(displayName="Table1", ref="A1:A2"))
                    sheet["C1"] = formula
                    workbook.save(path)
                    cached_workbook = openpyxl.load_workbook(path, data_only=True)
                    try:
                        self.assertIsNone(cached_workbook.active["C1"].value)
                    finally:
                        cached_workbook.close()
                    completed = subprocess.run(
                        [sys.executable, str(REPOSITORY_ROOT / ".agents/scripts/validate_financial_artifact.py"), str(path)],
                        capture_output=True,
                        text=True,
                    )
                    result = json.loads(completed.stdout)

                self.assertEqual(completed.returncode, expected_exit_code, completed.stdout + completed.stderr)
                self.assertEqual(result["status"], "fail" if expected_exit_code else "pass_with_limitations")
                self.assertEqual(result["structural_checks"]["formula_text"], "fail" if expected_exit_code else "pass")
                self.assertIn("FORMULA_EVALUATION_UNVERIFIED", result["limitations"])

class QuickLookValidationTests(unittest.TestCase):
    def test_timeout_returns_pass_with_full_render_limitation(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "document.docx"
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr(
                    "[Content_Types].xml",
                    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>',
                )
                archive.writestr(
                    "_rels/.rels",
                    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>',
                )
                archive.writestr(
                    "word/document.xml",
                    '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                    "<w:body><w:p/></w:body></w:document>",
                )

            with mock.patch.object(
                artifact_validator.shutil,
                "which",
                return_value="/usr/bin/qlmanage",
            ), mock.patch.object(
                artifact_validator.subprocess,
                "run",
                side_effect=subprocess.TimeoutExpired("qlmanage", 30),
            ):
                result, exit_code = artifact_validator.validate(path)

        self.assertEqual(exit_code, 0)
        self.assertEqual(result["status"], "pass_with_limitations")
        self.assertEqual(result["visual_review"]["status"], "unverified")
        self.assertEqual(result["visual_review"]["marker"], "FULL_RENDER_UNVERIFIED")
        self.assertIn("FULL_RENDER_UNVERIFIED", result["limitations"])


if __name__ == "__main__":
    unittest.main()
