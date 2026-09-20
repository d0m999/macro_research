#!/usr/bin/env python3
"""Validate DCF workbook structure and hard mathematical invariants."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any


EXCEL_ERRORS = {"#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A"}


def formula_has_unquoted_excel_error(formula_text: str) -> bool:
    """Return whether a formula contains an Excel error outside string literals."""
    unquoted: list[str] = []
    in_string = False
    string_start: int | None = None
    index = 0
    while index < len(formula_text):
        character = formula_text[index]
        if character != '"':
            if not in_string:
                unquoted.append(character)
            index += 1
            continue
        if in_string and index + 1 < len(formula_text) and formula_text[index + 1] == '"':
            index += 2
            continue
        in_string = not in_string
        string_start = index if in_string else None
        index += 1
    if string_start is not None:
        unquoted.extend(formula_text[string_start:])
    searchable_text = "".join(unquoted)
    return any(token in searchable_text for token in EXCEL_ERRORS)


def adjacent_number(sheet: Any, row: int, column: int) -> float | None:
    for offset in range(1, 5):
        value = sheet.cell(row, column + offset).value
        if isinstance(value, (int, float)):
            return float(value)
    return None


def find_labeled_number(sheet: Any, required_terms: tuple[str, ...]) -> float | None:
    for row in sheet.iter_rows(max_row=300, max_col=30):
        for cell in row:
            if not isinstance(cell.value, str):
                continue
            label = cell.value.lower()
            if all(term in label for term in required_terms):
                value = adjacent_number(sheet, cell.row, cell.column)
                if value is not None:
                    return value
    return None


def validate(path: Path, args: argparse.Namespace) -> dict[str, Any]:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for this optional DCF validator") from exc

    formulas = openpyxl.load_workbook(path, data_only=False, read_only=False)
    values = openpyxl.load_workbook(path, data_only=True, read_only=False)
    errors: list[str] = []
    warnings: list[str] = []
    checks: dict[str, Any] = {}

    checks["sheet_names"] = formulas.sheetnames
    if "DCF" not in formulas.sheetnames:
        errors.append("Required DCF sheet is missing")
    if "WACC" not in formulas.sheetnames:
        warnings.append("WACC sheet is absent; WACC may be embedded in the DCF sheet")

    formula_count = 0
    for sheet_name in formulas.sheetnames:
        formula_sheet = formulas[sheet_name]
        value_sheet = values[sheet_name]
        for row in formula_sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_count += 1
                    if formula_has_unquoted_excel_error(cell.value):
                        errors.append(f"Invalid reference literal in {sheet_name}!{cell.coordinate}")
                cached = value_sheet[cell.coordinate].value
                if isinstance(cached, str) and cached in EXCEL_ERRORS:
                    errors.append(f"Cached Excel error {cached} in {sheet_name}!{cell.coordinate}")
    checks["formula_count"] = formula_count
    if formula_count == 0:
        warnings.append("No workbook formulas were found")

    if "DCF" in values.sheetnames:
        dcf = values["DCF"]
        terminal_growth = find_labeled_number(dcf, ("terminal", "growth"))
        wacc_sheet = values["WACC"] if "WACC" in values.sheetnames else dcf
        wacc = find_labeled_number(wacc_sheet, ("wacc",))
        checks["terminal_growth"] = terminal_growth
        checks["wacc"] = wacc
        if terminal_growth is not None and wacc is not None and terminal_growth >= wacc:
            errors.append("Terminal growth must be lower than WACC for a Gordon Growth terminal value")
        elif terminal_growth is None or wacc is None:
            warnings.append("Terminal growth and WACC could not both be located from cached values")

        if wacc is not None and args.wacc_min is not None and wacc < args.wacc_min:
            warnings.append(f"WACC is below the user-supplied review bound {args.wacc_min:.4f}")
        if wacc is not None and args.wacc_max is not None and wacc > args.wacc_max:
            warnings.append(f"WACC is above the user-supplied review bound {args.wacc_max:.4f}")

        terminal_pv = find_labeled_number(dcf, ("terminal", "value", "pv"))
        enterprise_value = find_labeled_number(dcf, ("enterprise", "value"))
        proportion = None
        if terminal_pv is not None and enterprise_value not in (None, 0):
            proportion = terminal_pv / enterprise_value
        checks["terminal_value_proportion"] = proportion
        if proportion is not None and args.terminal_value_min is not None and proportion < args.terminal_value_min:
            warnings.append(
                "Terminal value proportion is below the user-supplied review bound "
                f"{args.terminal_value_min:.4f}"
            )
        if proportion is not None and args.terminal_value_max is not None and proportion > args.terminal_value_max:
            warnings.append(
                "Terminal value proportion is above the user-supplied review bound "
                f"{args.terminal_value_max:.4f}"
            )

    return {
        "file": str(path),
        "status": "fail" if errors else "pass",
        "structural_checks": checks,
        "errors": errors,
        "warnings": warnings,
        "limitations": [
            "Cached values may be stale; run the project artifact validator and disclose FORMULA_EVALUATION_UNVERIFIED when applicable."
        ],
    }


def parser() -> argparse.ArgumentParser:
    cli = argparse.ArgumentParser(
        description="Check DCF structure and hard math; heuristic bounds are optional user-supplied reviews."
    )
    cli.add_argument("file", type=Path)
    cli.add_argument("--wacc-min", type=float)
    cli.add_argument("--wacc-max", type=float)
    cli.add_argument("--terminal-value-min", type=float)
    cli.add_argument("--terminal-value-max", type=float)
    return cli


def main() -> int:
    args = parser().parse_args()
    if not args.file.is_file():
        print(json.dumps({"status": "fail", "error": f"File not found: {args.file}"}, indent=2))
        return 2
    try:
        result = validate(args.file, args)
    except Exception as exc:
        print(json.dumps({"status": "fail", "error": str(exc)}, indent=2))
        return 1
    print(json.dumps(result, indent=2))
    return 0 if result["status"] == "pass" else 1


if __name__ == "__main__":
    sys.exit(main())
